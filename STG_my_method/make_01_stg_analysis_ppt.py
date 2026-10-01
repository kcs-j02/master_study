from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt


OUT = "01_stg_analysis_summary.pptx"
WIDE = (13.333, 7.5)
NAVY = RGBColor(22, 43, 70)
BLUE = RGBColor(42, 104, 169)
TEAL = RGBColor(0, 139, 139)
ORANGE = RGBColor(222, 119, 42)
INK = RGBColor(35, 43, 52)
MUTED = RGBColor(91, 105, 119)
PALE = RGBColor(241, 246, 250)
WHITE = RGBColor(255, 255, 255)
LINE = RGBColor(205, 216, 225)


prs = Presentation()
prs.slide_width = Inches(WIDE[0])
prs.slide_height = Inches(WIDE[1])


def textbox(slide, x, y, w, h, text, size=20, color=INK, bold=False,
            align=PP_ALIGN.LEFT, font="Noto Sans CJK JP"):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = shape.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = Inches(0.06)
    tf.margin_right = Inches(0.06)
    tf.margin_top = Inches(0.03)
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    return shape


def rect(slide, x, y, w, h, fill, line=LINE, radius=False):
    kind = MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE if radius else MSO_AUTO_SHAPE_TYPE.RECTANGLE
    shape = slide.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.color.rgb = line
    shape.line.width = Pt(1)
    return shape


def title(slide, label, heading, subtitle=None):
    rect(slide, 0, 0, WIDE[0], 0.18, BLUE, BLUE)
    textbox(slide, 0.55, 0.42, 1.3, 0.32, label.upper(), 10, TEAL, True)
    textbox(slide, 0.55, 0.75, 12.1, 0.6, heading, 27, NAVY, True)
    if subtitle:
        textbox(slide, 0.58, 1.38, 11.9, 0.35, subtitle, 12, MUTED)


def footer(slide, number):
    textbox(slide, 0.55, 7.12, 8, 0.18, "STG提案手法 / Stage 1", 9, MUTED)
    textbox(slide, 12.1, 7.12, 0.65, 0.18, str(number), 9, MUTED, align=PP_ALIGN.RIGHT)


def bullet(slide, x, y, w, text, color=INK, size=17, accent=TEAL):
    rect(slide, x, y + 0.13, 0.09, 0.09, accent, accent, True)
    textbox(slide, x + 0.18, y, w - 0.18, 0.38, text, size, color)


def add_slide(label, heading, subtitle=None):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = WHITE
    title(slide, label, heading, subtitle)
    footer(slide, len(prs.slides))
    return slide


def arrow(slide, x1, y1, x2, y2, color=BLUE):
    line = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    line.line.color.rgb = color
    line.line.width = Pt(2)
    line.line.end_arrowhead = True


# 1. Cover
slide = prs.slides.add_slide(prs.slide_layouts[6])
slide.background.fill.solid()
slide.background.fill.fore_color.rgb = NAVY
rect(slide, 0, 0, 13.333, 0.2, TEAL, TEAL)
textbox(slide, 0.75, 1.0, 1.8, 0.35, "STAGE 1", 13, RGBColor(125, 223, 211), True)
textbox(slide, 0.75, 1.65, 11.8, 1.0, "STG解析", 38, WHITE, True)
textbox(slide, 0.78, 2.8, 10.8, 0.55, "タスク依存関係と並列性を、後段のStream設計へ渡せる形に変換", 19, RGBColor(219, 231, 240))
rect(slide, 0.78, 4.45, 11.8, 1.25, RGBColor(31, 61, 93), RGBColor(64, 100, 135), True)
textbox(slide, 1.08, 4.68, 11.2, 0.72, "STG  →  TaskSpec  →  TaskDFG  →  TaskLevels", 24, WHITE, True, PP_ALIGN.CENTER)
textbox(slide, 0.78, 6.62, 6, 0.25, "01_stg_analysis.hpp", 11, RGBColor(174, 194, 211))
textbox(slide, 12.1, 7.12, 0.65, 0.18, "1", 9, RGBColor(174, 194, 211), align=PP_ALIGN.RIGHT)

# 2. Objective
slide = add_slide("01 / 目的", "何を解析し、何を後段へ渡すか", "Stage 1の出力は、実行構成選択とStream割当の入力になる")
rect(slide, 0.65, 2.05, 3.65, 3.8, PALE, LINE, True)
textbox(slide, 0.95, 2.32, 3.0, 0.45, "入力", 16, BLUE, True)
textbox(slide, 0.95, 2.9, 3.0, 0.8, "STGファイル", 25, NAVY, True)
textbox(slide, 0.95, 3.78, 3.0, 0.7, "タスク / proc_time / preds", 15, MUTED)
rect(slide, 4.8, 2.05, 3.65, 3.8, RGBColor(231, 247, 244), RGBColor(157, 218, 208), True)
textbox(slide, 5.1, 2.32, 3.0, 0.45, "解析", 16, TEAL, True)
textbox(slide, 5.1, 2.9, 3.0, 0.8, "依存関係を構造化", 22, NAVY, True)
textbox(slide, 5.1, 3.78, 3.0, 0.7, "DFG構築 / level化 / 幅の抽出", 15, MUTED)
rect(slide, 8.95, 2.05, 3.75, 3.8, RGBColor(255, 244, 232), RGBColor(240, 191, 147), True)
textbox(slide, 9.25, 2.32, 3.0, 0.45, "出力", 16, ORANGE, True)
textbox(slide, 9.25, 2.9, 3.0, 0.8, "StgAnalysisResult", 22, NAVY, True)
textbox(slide, 9.25, 3.78, 3.0, 0.95, "graph / tasks / dfg / levels\nStream数決定に利用", 15, MUTED)
arrow(slide, 4.3, 3.95, 4.75, 3.95)
arrow(slide, 8.45, 3.95, 8.9, 3.95)

# 3. Flow
slide = add_slide("02 / 処理フロー", "解析は4段階で進む", "analyze_stg(path) がStage 1の入口")
steps = [
    ("01", "STG読込", "load_stg_without_comm", BLUE),
    ("02", "TaskSpec生成", "proc_timeで属性付与", TEAL),
    ("03", "DFG構築", "preds / succs / indeg", ORANGE),
    ("04", "Level化", "cycle検出 + 並列幅", RGBColor(128, 92, 165)),
]
for i, (num, name, detail, color) in enumerate(steps):
    x = 0.65 + i * 3.15
    rect(slide, x, 2.3, 2.55, 2.5, WHITE, color, True)
    rect(slide, x, 2.3, 2.55, 0.48, color, color, True)
    textbox(slide, x + 0.16, 2.39, 0.45, 0.22, num, 12, WHITE, True)
    textbox(slide, x + 0.18, 3.0, 2.15, 0.55, name, 21, NAVY, True)
    textbox(slide, x + 0.18, 3.75, 2.15, 0.55, detail, 13, MUTED)
    if i < 3:
        arrow(slide, x + 2.6, 3.55, x + 3.05, 3.55)
rect(slide, 1.0, 5.45, 11.2, 0.8, PALE, LINE, True)
textbox(slide, 1.25, 5.63, 10.7, 0.42, "戻り値: graph + tasks + dfg + levels  |  後段はlevelsから最大並列幅とStream数を決定", 15, NAVY, True, PP_ALIGN.CENTER)

# 4. Data structures
slide = add_slide("03 / データ構造", "タスクを「依存グラフ」と「実行レベル」で二重に表現", "局所的な隣接関係と、全体の並列性を分けて扱う")
rect(slide, 0.65, 2.0, 5.8, 4.35, PALE, LINE, True)
textbox(slide, 0.95, 2.25, 5.1, 0.4, "NodeInfo / TaskDFG", 20, NAVY, True)
for i, text in enumerate(["id: タスク識別子", "preds / succs: 前後関係", "indeg: 未処理の先行数", "proc_time: 実行時間", "order: 入力順を保持"]):
    bullet(slide, 1.0, 2.95 + i * 0.52, 4.9, text, size=15, accent=BLUE)
rect(slide, 6.9, 2.0, 5.8, 4.35, RGBColor(231, 247, 244), RGBColor(157, 218, 208), True)
textbox(slide, 7.2, 2.25, 5.1, 0.4, "TaskLevels", 20, NAVY, True)
textbox(slide, 7.2, 2.88, 5.0, 0.45, "vector<vector<int>>", 15, TEAL, True)
textbox(slide, 7.2, 3.45, 5.0, 0.55, "同じlevel内のタスクは、\n依存関係上は同時に開始可能", 17, INK)
for i, (label, tasks) in enumerate([("L0", "0"), ("L1", "1, 2"), ("L2", "3")]):
    x = 7.3 + i * 1.55
    rect(slide, x, 4.55, 1.25, 0.95, WHITE, TEAL, True)
    textbox(slide, x, 4.68, 1.25, 0.22, label, 12, TEAL, True, PP_ALIGN.CENTER)
    textbox(slide, x, 4.98, 1.25, 0.25, tasks, 18, NAVY, True, PP_ALIGN.CENTER)

# 5. Levelization
slide = add_slide("04 / Level化", "indegreeを減らしながら、トポロジカルに層を作る", "cycleがある場合は、levelを作れず例外として検出")
textbox(slide, 0.75, 2.02, 5.7, 0.4, "アルゴリズム", 17, BLUE, True)
for i, text in enumerate([
    "1. remaining_indeg = 各ノードのindegで初期化",
    "2. remaining_indeg == 0 の未割当タスクを集める",
    "3. そのタスクを同じlevelへ登録",
    "4. 後続タスクのindegを1つ減らす",
    "5. 全ノード割当まで繰り返す",
]):
    bullet(slide, 0.82, 2.55 + i * 0.52, 5.6, text, size=14, accent=BLUE)
rect(slide, 7.0, 2.0, 5.6, 4.1, WHITE, LINE, True)
textbox(slide, 7.3, 2.25, 4.8, 0.35, "例: 0 → {1,2} → 3", 17, NAVY, True)
nodes = [("0", 7.45, 3.2, BLUE), ("1", 8.65, 4.4, TEAL), ("2", 10.35, 4.4, TEAL), ("3", 11.3, 3.2, ORANGE)]
for label, x, y, color in nodes:
    rect(slide, x, y, 0.7, 0.7, color, color, True)
    textbox(slide, x, y + 0.13, 0.7, 0.3, label, 18, WHITE, True, PP_ALIGN.CENTER)
arrow(slide, 8.15, 3.55, 8.6, 4.62, MUTED)
arrow(slide, 8.15, 3.55, 10.3, 4.62, MUTED)
arrow(slide, 9.35, 4.75, 11.2, 3.55, MUTED)
arrow(slide, 10.95, 4.75, 11.2, 3.55, MUTED)
for x, label in [(7.36, "L0"), (8.55, "L1"), (11.2, "L2")]:
    textbox(slide, x, 5.43, 1.0, 0.25, label, 11, MUTED, True, PP_ALIGN.CENTER)
rect(slide, 7.3, 5.78, 4.8, 0.55, RGBColor(255, 244, 232), RGBColor(240, 191, 147), True)
textbox(slide, 7.5, 5.9, 4.4, 0.25, "level.empty() → cycle detected", 13, ORANGE, True, PP_ALIGN.CENTER)

# 6. Task classification and stream count
slide = add_slide("05 / 並列性の抽出", "proc_timeでタスク属性を付け、level幅からStream数を決める", "解析結果は後段の実行構成候補の上限を決める")
rect(slide, 0.65, 2.0, 5.9, 3.9, PALE, LINE, True)
textbox(slide, 0.95, 2.25, 5.1, 0.35, "タスク属性", 18, NAVY, True)
rect(slide, 1.0, 2.95, 2.3, 1.0, RGBColor(255, 244, 232), RGBColor(240, 191, 147), True)
textbox(slide, 1.0, 3.12, 2.3, 0.25, "proc_time > threshold", 12, ORANGE, True, PP_ALIGN.CENTER)
textbox(slide, 1.0, 3.49, 2.3, 0.25, "HEAVY / work_units", 14, NAVY, True, PP_ALIGN.CENTER)
rect(slide, 3.8, 2.95, 2.3, 1.0, RGBColor(231, 247, 244), RGBColor(157, 218, 208), True)
textbox(slide, 3.8, 3.12, 2.3, 0.25, "proc_time <= threshold", 12, TEAL, True, PP_ALIGN.CENTER)
textbox(slide, 3.8, 3.49, 2.3, 0.25, "LIGHT / work_units", 14, NAVY, True, PP_ALIGN.CENTER)
textbox(slide, 1.0, 4.65, 5.15, 0.7, "共通設定: parallel_sm_limit = 114\nproc_timeの基準は64 SM", 14, MUTED)
rect(slide, 6.95, 2.0, 5.75, 3.9, RGBColor(231, 247, 244), RGBColor(157, 218, 208), True)
textbox(slide, 7.25, 2.25, 5.1, 0.35, "Stream数", 18, NAVY, True)
textbox(slide, 7.25, 3.0, 5.1, 0.55, "S = min(max level width,\n    max_stream_count)", 23, TEAL, True, PP_ALIGN.CENTER)
textbox(slide, 7.25, 4.35, 5.1, 0.7, "最大5 Streamまで。\nlevel幅が並列性の上限になる。", 15, MUTED, align=PP_ALIGN.CENTER)

# 7. Summary
slide = add_slide("06 / まとめ", "Stage 1の要点", "後段の設計判断に必要な情報を、検証付きで一度に整える")
items = [
    ("構造化", "STGのpredsを使って、タスクの前後関係をTaskDFGへ変換"),
    ("並列性", "トポロジカルなlevelを作り、同時実行可能な幅を抽出"),
    ("分類", "proc_timeに基づきHEAVY / LIGHTとwork_unitsを設定"),
    ("安全性", "重複ID・未知の先行タスク・cycle・不正indegreeを検出"),
]
for i, (head, body) in enumerate(items):
    y = 1.95 + i * 1.08
    color = [BLUE, TEAL, ORANGE, RGBColor(128, 92, 165)][i]
    rect(slide, 0.75, y, 1.65, 0.7, color, color, True)
    textbox(slide, 0.75, y + 0.18, 1.65, 0.25, head, 16, WHITE, True, PP_ALIGN.CENTER)
    textbox(slide, 2.7, y + 0.08, 9.6, 0.52, body, 17, INK)
rect(slide, 0.75, 6.25, 11.85, 0.55, NAVY, NAVY, True)
textbox(slide, 1.0, 6.38, 11.35, 0.25, "出力: StgAnalysisResult  →  Stage 2以降のタスク重要度・Stream割当・実行構成選択", 14, WHITE, True, PP_ALIGN.CENTER)


prs.save(OUT)
print(OUT)